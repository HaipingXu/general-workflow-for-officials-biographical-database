#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
清理省长/书记总表中"多余的履历尾行"（事件行 / 公告碎片 / 重复 / 选举身份行）。

判定逻辑（每位官员 × 每张表，按 经历序号 排序；核心原则：删行必须有证据，
唯一的真任职行永不删）：

  R1 事件行    : 职务 命中"非任职事件"词（离休/退休/辞职/免职/开除/审查/查处/公诉/
                 宣判/不再担任/资格终止 …）→ 无条件删。
  R2 公告碎片  : source line(原文引用) 是"公告/叙述"类(免/辞/选举/当选/任命…为/决定/
                 批准/涉嫌/调查/查处/逮捕/公诉/宣判)，且本行起止退化成"点状/空"或时间错位，
                 且同单位另有真任职行 → 删（该行是公告被误抽成的碎片）。
  R3 完全重复  : 四字段(单位/职务/起/止)与更早行完全相同 → 删 经历序号 较大者。
  R4 结构化截断: 同单位且职务为对方子集、对方"至今"、本行退化(如 2022~2022) → 删截断那条。
  R5 选举/身份 : 职务∈委员/候补委员/代表 且 单位∈中央委员会/全国人大(排除政治局)，
                 且时间错位、source 为当选/选举 → 删（属选举身份，非任职）。

用法:
  uv run scripts/clean_trailing_rows.py                # dry-run，只生成待删清单
  uv run scripts/clean_trailing_rows.py --apply        # 备份→清理→保存→报告
  uv run scripts/clean_trailing_rows.py --apply --renumber-seq   # 同时把 经历序号 重排为连续

输出:
  temp/尾行清理_dryrun_<时间戳>.csv  (dry-run 待删清单)
  temp/尾行清理_applied_<时间戳>.csv (实删清单)
  temp/backups/尾行清理前_<时间戳>/  (--apply 时的原表备份)
"""
from __future__ import annotations
import argparse
import csv
import datetime as dt
import re
import shutil
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FILES = [
    ("省长", ROOT / "output/全省_省长数据库.xlsx", "省长数据库"),
    ("书记", ROOT / "output/全省_省委书记数据库.xlsx", "省委书记数据库"),
]

# ---- R1: 非任职"事件"词（出现在【职务】里）。已全表验证 0 误报。----
# 退休/离休 单独处理：排除"退休人员/退休金/离休干部局"等部门名词用法（防误删真职务）。
RETIRE = ("退休", "离休", "离职休养", "退二线")
RETIRE_NOUN = re.compile(r"(退休|离休)(人员|金|办|室|处|局|科|中心|管理|服务|干部|委员会|工作|证)")
EVENT_KW_OTHER = [
    "辞职", "辞去", "免职", "免去", "开除", "双开",
    "纪律审查", "监察调查", "审查调查", "立案", "被查处", "查处", "被查",
    "提起公诉", "公诉", "起诉", "宣判", "判处", "有期徒刑", "无期徒刑", "死刑",
    "受贿", "行贿", "撤职", "撤销党内", "撤销其", "罢免",
    "不再兼任", "不再担任", "资格终止", "终止职务", "逮捕", "取保候审", "执行逮捕",
]

# ---- R2/R5: source line 属"公告/叙述"类的判据 ----
ANNOUNCE_KW = re.compile(
    r"(免去|免职|不再担任|不再兼任|辞去|辞职|接受.{0,8}辞|的请求|"
    r"选举.{0,10}为|当选|增选|补选|任命.{0,10}为|提名|"
    r"中央决定|中共中央决定|中央批准|组织部|决定：|涉嫌|严重违纪|严重违法|"
    r"接受组织|接受.{0,6}审查|接受.{0,6}调查|被查处|查处|逮捕|提起公诉|"
    r"公诉|起诉|公开宣判|宣判|判处|开除|双开|资格终止|被免)"
)
# 公告/叙述常以"YYYY年M月[D日]，…"开头（结构化时间线则是 "YYYY.MM-" 区间）
ANNOUNCE_DATE = re.compile(r"^\s*(?:L\d+:\s*)?\d{4}年\d{1,2}月(?:\d{1,2}日)?")
# 结构化时间线行： "YYYY[.MM]-YYYY[.MM] …" 或 "YYYY[.MM]- …" 或 "YYYY-YYYY年 …"
STRUCT_LINE = re.compile(r"^\s*(?:L\d+:\s*)?\d{4}(?:\.\d{1,2})?\s*[-–—]")

# ---- R5: 选举/身份行 ----
MEMBER_JOB = re.compile(r"(候补委员|^委员$|^委员[、，]|人大代表|代表$)")
MEMBER_UNIT = re.compile(r"(中央委|中央委员会|全国人大|全国人民代表大会|党代|党的代表大会)")

# ---- M: 人工核定补删（通用规则因"单位串不一致"漏掉、但已逐条确认是重复/公告碎片）----
# 键 = (表tag, 省份, 姓名)，值 = {经历序号: 原因}
MANUAL_REMOVE = {
    ("省长", "吉林省", "高严"): {24: "M 人工核定: 「电力部」=「电力工业部」(seq20)叙述重复，单位串不一致漏删"},
    ("省长", "广东省", "孟凡利"): {41: "M 人工核定: 选举机关「广东省人大」被当单位，=省政府省长(seq37/40)，选举公告碎片"},
}


def colmap(ws):
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {n: i for i, n in enumerate(hdr)}, hdr


def year_of(v):
    if v is None:
        return None
    m = re.search(r"(\d{4})", str(v))
    return int(m.group(1)) if m else None


def tokens(job):
    s = str(job or "").strip()
    if not s or s == "None":
        return frozenset()
    parts = re.split(r"[、，,/／]", s)
    return frozenset(p.strip() for p in parts if p.strip())


def is_retire_event(job):
    j = str(job or "")
    if not any(w in j for w in RETIRE):
        return False
    return not RETIRE_NOUN.search(j)  # 排除部门名词用法


def is_event(job):
    j = str(job or "")
    return is_retire_event(j) or any(k in j for k in EVENT_KW_OTHER)


def comp_score(start, end):
    """完整度：起始有值+1，终止有值+1，'至今'额外+1（保证去重时保留最完整副本）。"""
    s = 1 if (start is not None and str(start) not in ("", "None")) else 0
    es = str(end or "")
    if "至今" in es:
        e = 2
    elif end is not None and es not in ("", "None"):
        e = 1
    else:
        e = 0
    return s + e


def is_announce(quote):
    q = str(quote or "")
    if ANNOUNCE_KW.search(q):
        return True
    # 以"YYYY年M月…"叙述开头，且不是结构化区间
    if ANNOUNCE_DATE.search(q) and not STRUCT_LINE.search(q):
        return True
    return False


def is_point_or_empty(start, end):
    """起止退化成"点状"(同值)或空。"""
    if start is None or end is None:
        return True
    se, ee = str(start), str(end)
    if se == "None" or ee == "None" or se == "" or ee == "":
        return True
    if start == end or se == ee:
        return True
    # 数值相等（如 2002 vs 2002.0）
    try:
        if float(start) == float(end):
            return True
    except (TypeError, ValueError):
        pass
    return False


def _dup_partner(rows, i, removed, *, same_unit=True):
    """找能"压制"第 i 行的同岗位副本：完整度更高，或同等完整但 seq 更小（保证保留最完整/最早副本）。"""
    r = rows[i]
    if not r["toks"]:
        return None
    for j, s in enumerate(rows):
        if j == i or j in removed:
            continue
        if same_unit and s["_u"] != r["_u"]:
            continue
        if not (r["toks"] <= s["toks"] or s["toks"] <= r["toks"]):
            continue
        if s["comp"] > r["comp"] or (s["comp"] == r["comp"] and j < i):
            return s
    return None


def analyze_person(rows, H):
    """rows: list of dict(seq, sheet_row, unit, job, start, end, quote). 返回 [(dict, reason)]。"""
    rows = sorted(rows, key=lambda r: (r["seq"] if isinstance(r["seq"], (int, float)) else 0))
    for r in rows:
        r["toks"] = tokens(r["job"])
        r["yr"] = year_of(r["start"])
        r["comp"] = comp_score(r["start"], r["end"])
        r["announce"] = is_announce(r["quote"])
        r["point"] = is_point_or_empty(r["start"], r["end"])
        r["_u"] = str(r["unit"])
    to_remove = []
    removed = set()  # 位置索引

    def max_year_before(i):
        ys = [rows[j]["yr"] for j in range(i) if j not in removed and rows[j]["yr"] is not None]
        return max(ys) if ys else None

    # ---- R1 事件行（全行扫描，无条件）----
    for i, r in enumerate(rows):
        if is_event(r["job"]):
            to_remove.append((r, "R1 事件行: 职务非任职事件"))
            removed.add(i)

    # ---- R3 完全重复（四字段一致，留 seq 最小）----
    seen = {}
    for i, r in enumerate(rows):
        if i in removed:
            continue
        key = (r["_u"], str(r["job"]), str(r["start"]), str(r["end"]))
        if key in seen:
            to_remove.append((r, f"R3 完全重复: =seq{rows[seen[key]]['seq']}"))
            removed.add(i)
        else:
            seen[key] = i

    # ---- R4 结构化截断重复（本行退化、同单位另有'至今'行覆盖同职务）----
    for i, r in enumerate(rows):
        if i in removed or not r["point"] or not r["toks"]:
            continue
        for j, s in enumerate(rows):
            if j == i or j in removed:
                continue
            if s["_u"] == r["_u"] and r["toks"] <= s["toks"] \
               and "至今" in str(s["end"]) and "至今" not in str(r["end"]):
                to_remove.append((r, f"R4 截断重复: 另有'至今'行 seq{s['seq']}"))
                removed.add(i)
                break

    # ---- R5 选举/身份行（错位 + 当选source + 委员/代表职务 + 中央委/全国人大单位，排除政治局）----
    for i, r in enumerate(rows):
        if i in removed or "政治局" in r["_u"]:
            continue
        if not (MEMBER_JOB.search(str(r["job"] or "")) and MEMBER_UNIT.search(r["_u"])):
            continue
        myr = max_year_before(i)
        if r["yr"] is not None and myr is not None and r["yr"] < myr and r["announce"]:
            to_remove.append((r, f"R5 选举/身份行: 错位({r['yr']}<{myr})"))
            removed.add(i)

    # ---- R2 公告碎片（source是公告/叙述 + 退化或错位 + 同单位另有更完整真任职行）----
    for i, r in enumerate(rows):
        if i in removed or not r["announce"]:
            continue
        myr = max_year_before(i)
        oo = r["yr"] is not None and myr is not None and r["yr"] < myr
        if not (r["point"] or oo):
            continue
        partner = _dup_partner(rows, i, removed, same_unit=True)
        if partner is not None:
            tag = "点状" if r["point"] else f"错位({r['yr']}<{myr})"
            to_remove.append((r, f"R2 公告碎片: {tag}, 留seq{partner['seq']}"))
            removed.add(i)

    # ---- R6 错位重复（错位 + 同单位/同职务、且【起止年份都相同】=同一段任职的复述）----
    # 严格要求 partner 与本行起止年份一致，避免把"相邻的不同任期"误判为重复（如卢东亮 seq17）。
    for i, r in enumerate(rows):
        if i in removed:
            continue
        myr = max_year_before(i)
        if not (r["yr"] is not None and myr is not None and r["yr"] < myr):
            continue
        r_end_yr = year_of(r["end"])
        partner = None
        for j, s in enumerate(rows):
            if j == i or j in removed or s["_u"] != r["_u"] or not r["toks"]:
                continue
            if not (r["toks"] <= s["toks"] or s["toks"] <= r["toks"]):
                continue
            if year_of(s["start"]) != r["yr"] or year_of(s["end"]) != r_end_yr:
                continue  # 必须是同一段任职（起止年份一致）才算复述
            if s["comp"] > r["comp"] or (s["comp"] == r["comp"] and j < i):
                partner = s
                break
        if partner is not None:
            to_remove.append((r, f"R6 错位重复: 错位({r['yr']}<{myr}), 留seq{partner['seq']}"))
            removed.add(i)

    return to_remove


def run(apply: bool, renumber: bool):
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    all_removals = []  # dict rows for csv
    per_file_remove = {}  # path -> set(sheet_row)

    for tag, path, sheet in FILES:
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        H, hdr = colmap(ws)
        people = defaultdict(list)
        for ridx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            people[(row[H["省份"]], row[H["姓名"]])].append({
                "sheet_row": ridx,
                "seq": row[H["经历序号"]],
                "unit": row[H["供职单位"]],
                "job": row[H["职务"]],
                "start": row[H["起始时间"]],
                "end": row[H["终止时间"]],
                "quote": row[H["原文引用"]],
            })
        remove_rows = set()
        for (prov, name), rows in people.items():
            results = analyze_person(rows, H)
            done_seq = {r["seq"] for r, _ in results}
            # M 人工核定补删（通用规则漏掉、已逐条确认）
            manual = MANUAL_REMOVE.get((tag, prov, name), {})
            for r in rows:
                if r["seq"] in manual and r["seq"] not in done_seq:
                    results.append((r, manual[r["seq"]]))
            for r, reason in results:
                remove_rows.add(r["sheet_row"])
                all_removals.append({
                    "表": tag, "省份": prov, "姓名": name, "经历序号": r["seq"],
                    "起始": r["start"], "终止": r["end"],
                    "供职单位": r["unit"], "职务": r["job"],
                    "删除规则": reason,
                    "原文引用": str(r["quote"])[:120],
                })
        per_file_remove[(tag, path, sheet)] = remove_rows
        wb.close()

    # 排序输出（按 表/省/姓名/seq）
    all_removals.sort(key=lambda d: (d["表"], d["省份"], d["姓名"],
                                     d["经历序号"] if isinstance(d["经历序号"], (int, float)) else 0))
    mode = "applied" if apply else "dryrun"
    out_csv = ROOT / f"temp/尾行清理_{mode}_{ts}.csv"
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["表", "省份", "姓名", "经历序号", "起始", "终止",
                                          "供职单位", "职务", "删除规则", "原文引用"])
        w.writeheader()
        w.writerows(all_removals)

    # 汇总
    by_rule = defaultdict(int)
    for d in all_removals:
        by_rule[d["删除规则"].split(":")[0]] += 1
    print(f"\n{'='*60}\n{'实删' if apply else 'DRY-RUN'}  共 {len(all_removals)} 行")
    for k in sorted(by_rule):
        print(f"  {k}: {by_rule[k]} 行")
    for (tag, path, sheet), rm in per_file_remove.items():
        print(f"  [{tag}] {len(rm)} 行")
    print(f"清单: {out_csv}")

    if not apply:
        print("\n>> dry-run，未改动任何数据。确认后加 --apply 执行。")
        return

    # ---- 实删：备份 → 删行 → (可选)重排 seq → 保存 ----
    backup_dir = ROOT / f"temp/backups/尾行清理前_{ts}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    for tag, path, sheet in FILES:
        shutil.copy2(path, backup_dir / path.name)
    print(f"已备份原表 → {backup_dir}")

    for tag, path, sheet in FILES:
        rm = per_file_remove[(tag, path, sheet)]
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        H, hdr = colmap(ws)
        for ridx in sorted(rm, reverse=True):  # 从下往上删，避免行号错位
            ws.delete_rows(ridx, 1)
        if renumber:
            # 按 省份+姓名 把 经历序号 重排为 1..K
            seq_col = H["经历序号"] + 1
            prov_col = H["省份"] + 1
            name_col = H["姓名"] + 1
            counters = defaultdict(int)
            for ridx in range(2, ws.max_row + 1):
                key = (ws.cell(ridx, prov_col).value, ws.cell(ridx, name_col).value)
                counters[key] += 1
                ws.cell(ridx, seq_col).value = counters[key]
        wb.save(path)
        print(f"[{tag}] 删除 {len(rm)} 行 → 保存 {path.name}" + ("（已重排经历序号）" if renumber else ""))
    print(f"\n实删清单: {out_csv}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="实际写入（默认仅 dry-run）")
    ap.add_argument("--renumber-seq", action="store_true", help="实删后把经历序号重排为连续")
    args = ap.parse_args()
    run(apply=args.apply, renumber=args.renumber_seq)
