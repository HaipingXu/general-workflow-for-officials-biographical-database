#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
统一 起始时间/终止时间 列格式为文本 "YYYY.MM"：
  - 有月份 → "1998.07"；无月份 → "1998.00"
  - int 1963 → "1963.00"；float 2003.03 → "2003.03"；float 2023.1(=10月) → "2023.10"
  - "2023.01.14"(带日) → "2023.01"；"至今"/空 → 不变
  - 单元格设为文本格式(@)，避免 Excel 重新解释成数字/日期

依据: float 为"年+月/100"编码，:.2f 可准确还原月份（含掉零的 10 月），已对照原文引用验证。

用法:
  uv run scripts/normalize_time_format.py           # 报告模式（不写）
  uv run scripts/normalize_time_format.py --apply    # 写入（建议先备份 temp/backups/）
"""
from __future__ import annotations
import argparse
import datetime as dt
import re
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FILES = [
    (ROOT / "output/全省_省长数据库.xlsx", "省长数据库"),
    (ROOT / "output/全省_省委书记数据库.xlsx", "省委书记数据库"),
]
COLS = ["起始时间", "终止时间"]

_STR_YM = re.compile(r"^(\d{4})\.(\d{1,2})(?:\.\d{1,2})?$")  # YYYY.MM 或 YYYY.MM.DD
_STR_Y = re.compile(r"^(\d{4})$")                              # YYYY


def normalize(v):
    """返回 (新值, 是否改动, 异常说明or None)。"""
    if v is None:
        return None, False, None
    if isinstance(v, bool):
        return v, False, f"布尔值?{v!r}"
    if isinstance(v, int):
        return f"{v}.00", True, None
    if isinstance(v, float):
        yr = int(v)
        mo = round((v - yr) * 100)
        if mo > 12 or mo < 0:
            return v, False, f"月份异常 {v!r}->{mo}"
        return f"{yr}.{mo:02d}", True, None
    if isinstance(v, dt.datetime):
        return f"{v.year}.{v.month:02d}", True, None
    s = str(v).strip()
    if s == "":
        return None, (v is not None), None
    if s == "至今":
        return "至今", (v != "至今"), None
    m = _STR_YM.match(s)
    if m:
        yr, mo = m.group(1), int(m.group(2))
        if mo > 12:
            return v, False, f"字符串月份异常 {s!r}"
        new = f"{yr}.{mo:02d}"
        return new, (new != s), None
    m = _STR_Y.match(s)
    if m:
        return f"{m.group(1)}.00", True, None
    return v, False, f"无法识别 {s!r}"


def run(apply: bool):
    for path, sheet in FILES:
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        H = {n: i for i, n in enumerate(hdr)}
        changed = 0
        anomalies = []
        before = Counter()
        for ridx in range(2, ws.max_row + 1):
            for name in COLS:
                cell = ws.cell(ridx, H[name] + 1)
                v = cell.value
                before[type(v).__name__] += 1
                new, ch, warn = normalize(v)
                if warn:
                    anomalies.append((ridx, name, repr(v), warn))
                if apply:
                    if ch:
                        cell.value = new
                    cell.number_format = "@"  # 文本，统一显示
                if ch:
                    changed += 1
        print(f"\n[{sheet}] 待改/已改 {changed} 个单元格；原类型分布: {dict(before)}")
        if anomalies:
            print(f"  ⚠ 异常 {len(anomalies)} 处:")
            for a in anomalies[:30]:
                print("    ", a)
        else:
            print("  ✓ 无异常")
        if apply:
            wb.save(path)
            print(f"  → 已保存 {path.name}")
        wb.close()
    if not apply:
        print("\n>> 报告模式，未写入。确认无异常后加 --apply。")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    run(apply=ap.parse_args().apply)
