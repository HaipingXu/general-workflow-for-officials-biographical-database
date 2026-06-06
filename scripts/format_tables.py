#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
给全省总表做可读性排版：自适应列宽 + 蓝色表头(白色加粗) + 冻结首行 + 自动筛选。

用法:  uv run scripts/format_tables.py
注意:  只改格式不改数据。建议先备份（temp/backups/）。
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
FILES = [
    (ROOT / "output/全省_省长数据库.xlsx", "省长数据库"),
    (ROOT / "output/全省_省委书记数据库.xlsx", "省委书记数据库"),
]

HEADER_FILL = PatternFill("solid", fgColor="4472C4")   # 蓝色
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)  # 白色加粗
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

# 每列宽度上限（按内容类型）：长文本列给大些，其余收紧
CAP_WIDE = {"原文引用": 60}
CAP_MED = {"judge1con": 45, "judge2con": 45, "judge3con": 45, "judge4con": 45,
           "落马原因": 40, "备注栏": 40, "供职单位": 26, "职务": 26}
CAP_DEFAULT = 22
MINW = 6


def disp_len(s):
    """显示宽度：CJK 记 2，其余记 1。"""
    s = "" if s is None else str(s)
    return sum(2 if ord(c) > 0x2E80 else 1 for c in s)


def run():
    for path, sheet in FILES:
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        ncol = len(headers)

        # 列宽：取 表头与该列数据的最大显示宽度，按列类型封顶
        maxw = [disp_len(h) for h in headers]
        for row in ws.iter_rows(min_row=2, values_only=True):
            for j in range(ncol):
                w = disp_len(row[j])
                if w > maxw[j]:
                    maxw[j] = w
        for j, name in enumerate(headers):
            cap = CAP_WIDE.get(name) or CAP_MED.get(name) or CAP_DEFAULT
            width = max(MINW, min(maxw[j] + 2, cap))
            ws.column_dimensions[get_column_letter(j + 1)].width = width

        # 表头：蓝底白字加粗居中
        for c in ws[1]:
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = HEADER_ALIGN
        ws.row_dimensions[1].height = 22

        # 冻结首行 + 自动筛选
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        wb.save(path)
        print(f"[{sheet}] 列宽/蓝色表头/冻结/筛选 已应用 → {path.name}")


if __name__ == "__main__":
    run()
